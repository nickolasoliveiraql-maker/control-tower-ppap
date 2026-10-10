import pandas as pd
import plotly.express as px
import streamlit as st
import openpyxl
import os
import warnings

warnings.filterwarnings('ignore', category=UserWarning, module='openpyxl')
st.set_page_config(page_title="Control Tower - Matriz de Risco", layout="wide")

# ==============================================================================
# CÉREBRO RASTREADOR DE DADOS
# ==============================================================================
@st.cache_data(ttl=300)
def ler_excel_avancado(caminho_arquivo):
    if not os.path.exists(caminho_arquivo): return None
    wb = openpyxl.load_workbook(caminho_arquivo, data_only=True)
    sheet = wb["COMPONENTES ATIVOS"] if "COMPONENTES ATIVOS" in wb.sheetnames else wb.worksheets

    linha_header = 1
    for r in range(1, 6):
        for c in range(1, 4):
            val = str(sheet.cell(row=r, column=c).value).upper()
            if "CÓDIGO" in val or "CODIGO" in val:
                linha_header = r
                break
        if linha_header == r: break

    mapa = {'CODIGO': -1, 'GM': -1, 'VW': -1, 'FORN': -1, 'STATUS': -1, 'FAMILIA': -1, 'DIAS_ATRASO': -1}
    for col_idx in range(1, sheet.max_column + 1): 
        cell_val = str(sheet.cell(row=linha_header, column=col_idx).value).upper().strip()
        if cell_val in ["NONE", ""]: continue
        idx_real = col_idx - 1
        if "CÓDIGO DO COMPONENTE" in cell_val or "CODIGO DO COMPONENTE" in cell_val: mapa['CODIGO'] = idx_real
        elif cell_val == "GM" or "PN GM" in cell_val: mapa['GM'] = idx_real
        elif cell_val == "VW" or "PN VW" in cell_val: mapa['VW'] = idx_real
        elif "NOME FORNECEDOR" in cell_val or "FORNECEDOR" in cell_val: mapa['FORN'] = idx_real
        elif "STATUS (FÓRMULA)" in cell_val or cell_val == "STATUS":
            if mapa['STATUS'] == -1: mapa['STATUS'] = idx_real
        elif "DESCRIÇÃO SIMPLES" in cell_val or "DESCRICAO SIMPLES" in cell_val: mapa['FAMILIA'] = idx_real
        elif "QTD DIAS EM RELAÇÃO AO STATUS" in cell_val or "DIAS EM RELACAO" in cell_val: mapa['DIAS_ATRASO'] = idx_real

    dados = []
    def limpar(val): 
        if val is None: return ""
        txt = str(val).strip().upper()
        return "" if txt in ["NONE", "NAN", "NA", "X", "-"] else txt
        
    def converter_inteiro(val):
        try: 
            if val is None: return 0
            txt = str(val).strip().upper()
            if txt in ["", "NA", "NAN", "NONE", "X", "-"]: return 0
            return int(float(txt))
        except: return 0

    for row_idx in range(linha_header + 1, sheet.max_row + 1):
        row = [sheet.cell(row=row_idx, column=c).value for c in range(1, sheet.max_column + 1)]
        if len(row) <= max(mapa.values()): continue
        codigo = limpar(row[mapa['CODIGO']])
        if codigo == "" or "CÓDIGO" in codigo: continue
        
        dados.append({
            'Codigo': codigo, 'GM': limpar(row[mapa['GM']]), 'VW': limpar(row[mapa['VW']]),
            'Fornecedor': "NÃO INFORMADO" if limpar(row[mapa['FORN']]) in ["", "0"] else limpar(row[mapa['FORN']]), 
            'Status': "OK" if limpar(row[mapa['STATUS']]) in ["", "0"] else limpar(row[mapa['STATUS']]), 
            'Familia': "Outros" if limpar(row[mapa['FAMILIA']]) in ["", "0"] else str(row[mapa['FAMILIA']]).strip().title(), 
            'DiasAtraso': converter_inteiro(row[mapa['DIAS_ATRASO']])
        })
    wb.close()
    df = pd.DataFrame(dados)
    df['Cliente'] = df.apply(lambda r: 'COMUM (GM/VW)' if r['GM']!="" and r['VW']!="" else ('GM' if r['GM']!="" else ('VW' if r['VW']!="" else 'OUTROS')), axis=1)
    return df

# ==============================================================================
# INTERFACE VISUAL DO DASHBOARD
# ==============================================================================
st.title("📊 CONTROL TOWER - MATRIZ DE RISCO")
st.sidebar.header("📁 Carga de Dados")
arquivo_carregado = st.sidebar.file_uploader("Arraste a planilha atualizada aqui", type=["xlsx", "xlsm"])

# Fallback para leitura local caso o arquivo seja colocado na mesma pasta
caminho_final = "INSP_TRABALHAR.xlsx" if os.path.exists("INSP_TRABALHAR.xlsx") else ("INSP_TRABALHAR.xlsm" if os.path.exists("INSP_TRABALHAR.xlsm") else None)
if arquivo_carregado is not None:
    with open("temp_planilha.xlsx", "wb") as f: f.write(arquivo_carregado.getbuffer())
    caminho_final = "temp_planilha.xlsx"

df_master = ler_excel_avancado(caminho_final) if caminho_final else None

if df_master is not None:
    st.sidebar.header("🔍 Filtros Operacionais")
    filtro_cliente = st.sidebar.multiselect("Montadora", options=df_master['Cliente'].unique())
    filtro_forn = st.sidebar.multiselect("Fornecedor", options=sorted(df_master['Fornecedor'].unique()))

    dff = df_master.copy()
    if filtro_cliente: dff = dff[dff['Cliente'].isin(filtro_cliente)]
    if filtro_forn: dff = dff[dff['Fornecedor'].isin(filtro_forn)]

    atrasados = dff[dff['Status'].str.contains('ATRASADO|ATENÇÃO|ATENCAO', regex=True, na=False)]
    dias_validos = atrasados[atrasados['DiasAtraso'] > 0]['DiasAtraso']
    media_aging = f"{int(dias_validos.mean())} dias" if not dias_validos.empty else "0 dias"

    col1, col2, col3 = st.columns(3)
    col1.metric("TOTAL ATIVOS", len(dff))
    col2.metric("PPAPs ATRASADOS", len(atrasados))
    col3.metric("DIAS EM MÉDIA DE ATRASO", media_aging)

    st.markdown("---")
    g1, g2 = st.columns(2)
    g1.plotly_chart(px.pie(dff, names='Status', hole=0.6, title="Distribuição Geral de Status", color_discrete_map={'OK': '#00e676', 'ATENÇÃO': '#fbc02d', 'ATRASADO': '#ff4d4d', 'REPROVADO': '#d32f2f'}, template="plotly_dark"), use_container_width=True)
    g2.plotly_chart(px.bar(dff['Familia'].value_counts().reset_index(), x='Familia', y='count', title="Volume por Família de Componente", template="plotly_dark", color_discrete_sequence=['#66fcf1']), use_container_width=True)

    g3, g4 = st.columns([1.5, 1])
    df_forn_atraso = atrasados.groupby('Fornecedor')['Codigo'].count().reset_index()
    if not df_forn_atraso.empty:
        fig_gargalo = px.bar(df_forn_atraso.sort_values(by='Codigo', ascending=True).tail(10), x='Codigo', y='Fornecedor', orientation='h', title="Top 10 Fornecedores com Mais Códigos Atrasados", labels={'Codigo': 'Qtd de Códigos Atrasados'}, template="plotly_dark", color_discrete_sequence=['#ff4d4d'])
        fig_gargalo.update_layout(yaxis={'type': 'category'})
        g3.plotly_chart(fig_gargalo, use_container_width=True)
    else:
        g3.plotly_chart(px.bar(title="Nenhum atraso registrado no filtro ativo", template="plotly_dark"), use_container_width=True)
    g4.plotly_chart(px.pie(dff, names='Cliente', title="Dependência por Cliente (Montadora)", color_discrete_sequence=['#45f3ff', '#ff007f', '#ffffff', '#7928ca'], template="plotly_dark"), use_container_width=True)
else:
    st.info("💡 Por favor, use o botão na barra lateral esquerda para carregar a sua planilha `INSP_TRABALHAR` e ativar os gráficos.")
